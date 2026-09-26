import os
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def generar_word(datos, ruta_guardado):
    doc = Document()

    # Ajuste de márgenes
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)

    estilo_normal = doc.styles['Normal']
    fuente = estilo_normal.font
    fuente.name = 'Calibri'
    fuente.size = Pt(11)

    # Cabecera
    p = doc.add_paragraph('')
    p.paragraph_format.space_after = Pt(0)
    p.add_run('Clave de documento: ').bold = True
    p.add_run('IN importación temporal de bienes')
    
    r = doc.add_paragraph('')
    r.paragraph_format.space_after = Pt(0)
    r.add_run('Régimen Anexo 22: ').bold = True
    r.add_run('ITE Temporales de importación')

    # Descripción de la mercancía
    doc.add_heading('Descripción de la mercancía', 1)
    desc = doc.add_paragraph(f"VARILLA DE ALUMINIO ENROLLADA {datos['varilla_aluminio']}, SIN RECUBRIMIENTO, SIN ACABADO, SIN ACCESORIOS INTEGRADOS, DIAMETRO {datos['diametro']} (INCHES), CON LOS SIGUIENTES PORCENTAJES DE COMPOSICION QUIMICA: ")

    # Composición Química (Bucle Dinámico por cada Heat)
    doc.add_heading('Composición Química', 1)

    for heat_id, heat_data in datos['heats'].items():
        p_heat = doc.add_paragraph('')
        p_heat.add_run(f"HEAT# {heat_id}\n").bold = True
        p_heat.add_run("MELT ORIGIN: CANADA").bold = True
        p_heat.paragraph_format.space_after = Pt(0)

        table = doc.add_table(rows=0, cols=2)
        table.autofit = False 
        table.allow_autofit = False

        for elemento, val in heat_data['quimicos'].items():
            row_cells = table.add_row().cells
            row_cells[0].text = str(elemento).capitalize()
            row_cells[1].text = str(val)
            row_cells[0].width = Cm(1.5)
            row_cells[1].width = Cm(2.0)
            for pf in row_cells[0].paragraphs: pf.paragraph_format.space_after = Pt(0)
            for pf in row_cells[1].paragraphs: pf.paragraph_format.space_after = Pt(0)

        # Salto pequeño entre Heats
        doc.add_paragraph() 

    # --- PÁRRAFO DE CANTIDADES ACTUALIZADO ---
    parrafo_cantidades = (
        f"LA CANTIDAD EXPRESADA EN EL CERTIFICADO NUMERO {datos['cert_num']}, CON TRIP NUMBER (NUMERO DE VIAJE) {datos['trip_num']}, "
        f"CONTEMPLA EL TOTAL DE {datos['total_libras']:,.2f} LIBRAS (PRECIO POR LIBRA {datos['precio_x_libra']:.3f} DOLARES) QUE CONVERTIDAS A KILOGRAMOS SON {datos['peso_kg']:,} KGS "
        f"(PRECIO UNITARIO POR KILOGRAMO DE {datos['precio_unit_kg']:.3f} DOLARES), VALOR TOTAL DEL PRODUCTO {datos['valor_total']:,.2f} DOLARES"
    )
    p_cant = doc.add_paragraph(parrafo_cantidades)
    p_cant.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Fracción
    doc.add_paragraph("LA FRACCION ARANCELARIA A UTILIZAR INCLUYENDO EL NICO ES XXXXXXXXXX")
    frac = doc.add_paragraph('')
    frac.add_run("FRACCION\n").bold = True
    frac.add_run("XXXXXXXX (xx)")

    # Cantidad
    cant = doc.add_paragraph('')
    cant.add_run("CANTIDAD\n").bold = True
    cant.add_run(f"{datos['peso_kg']:,} KGS ({datos['total_libras']:,.2f} LBS)")

    # Dolares
    usd = doc.add_paragraph('')
    usd.add_run("DOLARES\n").bold = True
    usd.add_run(f"{datos['valor_total']:,.2f}")

    # País de origen y exportador
    po = doc.add_paragraph('')
    po.add_run("PAIS ORIGEN: ").bold = True
    po.add_run("CANADA")

    pe = doc.add_paragraph('')
    pe.add_run("PAIS EXPORTADOR: ").bold = True
    pe.add_run("CANADA")

    # Certificado
    cert = doc.add_paragraph('')
    cert.add_run("CERTIFICADO\n").bold = True
    cert.add_run(f"{datos['cert_num']}, CON TRIP NUMBER (NUMERO DE VIAJE) {datos['trip_num']}.")

    # Fecha
    fecha = doc.add_paragraph('')
    fecha.add_run("FECHA\n").bold = True
    fecha.add_run(f"{datos['fecha']}")

    # --- OBSERVACIONES DINÁMICAS (Lógica de Gramática) ---
    texto_base_obs = (
        f"LA CANTIDAD EXPRESADA EN EL CERTIFICADO NUMERO {datos['cert_num']}, CON TRIP NUMBER (NUMERO DE VIAJE) {datos['trip_num']}, "
        f"CONTEMPLA EL TOTAL DE {datos['total_libras']:,.2f} LIBRAS QUE CONVERTIDAS A KILOGRAMOS SON {datos['peso_kg']:,} KGS.\n"
    )
    
    # Construimos la oración de "LOS CUALES SE INTEGRAN..."
    lista_pesos = []
    for heat_id, heat_data in datos['heats'].items():
        lbs_format = f"{heat_data['peso_lbs']:,.2f}"
        kgs_format = f"{int(heat_data['peso_kgs']):,}"
        lista_pesos.append(f"DEL HEAT {heat_id} EL PESO DE {lbs_format} LBS QUE CONVERTIDAS A KILOGRAMOS SON {kgs_format} KGS")
    
    if len(lista_pesos) > 1:
        # Une todo con comas y pone una "Y" antes del último
        oracion_integran = "LOS CUALES SE INTEGRAN " + ", ".join(lista_pesos[:-1]) + " Y " + lista_pesos[-1] + "."
    else:
        oracion_integran = "LOS CUALES SE INTEGRAN " + lista_pesos[0] + "."

    obs = doc.add_paragraph('')
    obs.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    obs.add_run('OBSERVACIONES\n').bold = True
    obs.add_run(texto_base_obs + oracion_integran)

    # INFORMACION DE LA EMPRESA
    emp = doc.add_paragraph('')
    emp.add_run("TIPO DE PERSONA\n").bold = True
    emp.add_run("MORAL")

    ds = doc.add_paragraph('')
    ds.add_run("DENOMINACION SOCIAL\n").bold = True
    ds.add_run(datos['denomi_social'].upper())

    empdir = doc.add_paragraph('')
    empdir.add_run("DIRECCION\n").bold = True
    empdir.add_run(datos['address'].upper())

    ef = doc.add_paragraph('')
    ef.add_run('ENTIDAD FEDERATIVA\n').bold = True
    ef.add_run("N/A")

    pf = doc.add_paragraph('')
    pf.add_run("REPRESENTACION FEDERAL\n").bold = True
    pf.add_run("N/A")

    # Guardamos usando el nombre del viaje (Trip)
    nombre_archivo = f"Plantilla_Importacion_{datos['trip_num']}.docx"
    ruta_completa = os.path.join(ruta_guardado, nombre_archivo)
    doc.save(ruta_completa)
    return nombre_archivo
    